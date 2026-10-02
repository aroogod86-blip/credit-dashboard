"""
포트 운용 탭 exporter
- 엑셀(포트운용_IG크레딧_*.xlsx)의 계산 결과를 읽어 대시보드용 HTML 조각을 만든다.
- 엑셀이 원천(source of truth). 이 스크립트는 값을 읽기만 하고 계산 로직은 갖지 않는다.
  (예외: 스왑분석은 대시보드에서 종목을 바꿔볼 수 있도록 [시장분석] 표를 내보내 브라우저에서 같은 규칙으로 필터링)

사용법
  1) 엑셀에서 원자료 붙여넣기 → 저장 (엑셀이 저장할 때 계산값이 같이 저장됨)
  2) python export_port_tab.py 포트운용_IG크레딧_2026-10-01.xlsx
       → port_tab_preview.html (단독 미리보기), port_tab.json 생성
  3) export_credit_dashboard.py 안에서:
       from export_port_tab import render_port_tab
       html_fragment = render_port_tab(xlsx_path)   # <section id="port-tab">…</section>
     를 기존 탭 컨테이너에 끼워 넣으면 됨. CSS는 #port-tab 아래로만 적용되어 기존 스타일과 충돌하지 않음.
"""
import json, sys, html, datetime as dt
from pathlib import Path
from openpyxl import load_workbook


# ---------------------------------------------------------------- 읽기
def _v(ws, ref):
    return ws[ref].value


def _num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def read_workbook(path):
    wb = load_workbook(path, data_only=True)
    S = wb["요약"]
    if _v(S, "B5") is None:
        raise RuntimeError(
            "엑셀에 계산값이 없습니다. 엑셀에서 파일을 열어 저장한 뒤 다시 실행하세요 "
            "(다른 도구로 저장하면 계산값이 비어 있을 수 있음).")

    out = {"generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
           "label": _v(S, "A2") or "", "warning": _v(S, "A3") or ""}

    # 요약 KPI (A5:B21)
    kpi = []
    for r in range(5, 40):
        k, v = S.cell(r, 1).value, S.cell(r, 2).value
        if k is None:
            break
        kpi.append({"k": k, "v": v})
    out["kpi"] = kpi

    # 시나리오 (D5:L12)
    scen_hdr = [S.cell(5, c).value for c in range(4, 13)]
    scen = []
    for r in range(6, 13):
        row = [S.cell(r, c).value for c in range(4, 13)]
        if row[0] is None:
            break
        scen.append(row)
    out["scen"] = {"hdr": scen_hdr, "rows": scen}
    out["breakeven"] = [{"k": S.cell(r, 4).value, "v": S.cell(r, 12).value} for r in (14, 15, 16)]

    # 운용코드별 (D19:K)
    fh = [S.cell(19, c).value for c in range(4, 12)]
    funds = []
    for r in range(20, 60):
        row = [S.cell(r, c).value for c in range(4, 12)]
        if row[0] is None:
            continue
        if not _num(row[1]):
            continue
        funds.append(row)
    out["funds"] = {"hdr": fh, "rows": funds}

    # 버킷분석: '■' 로 시작하는 블록
    B = wb["버킷분석"]
    buckets = []
    r = 1
    while r <= B.max_row:
        a = B.cell(r, 1).value
        if isinstance(a, str) and a.startswith("■"):
            name = a.lstrip("■ ").strip()
            hdr = [B.cell(r + 1, c).value for c in range(1, 21)]
            rows = []
            rr = r + 2
            while rr <= B.max_row and B.cell(rr, 1).value not in (None, ""):
                rows.append([B.cell(rr, c).value for c in range(1, 21)])
                if B.cell(rr, 1).value == "합계":
                    break
                rr += 1
            buckets.append({"name": name, "hdr": hdr, "rows": rows})
            r = rr + 1
        else:
            r += 1
    out["buckets"] = buckets

    # 매매후보 두 블록
    T = wb["매매후보"]
    cand = []
    for r in range(1, T.max_row + 1):
        a = T.cell(r, 1).value
        if isinstance(a, str) and (a.startswith("▲") or a.startswith("▼")):
            hdr = [T.cell(r + 1, c).value for c in range(1, 16)]
            rows = []
            for rr in range(r + 2, r + 22):
                row = [T.cell(rr, c).value for c in range(1, 16)]
                if row[1] in (None, ""):
                    continue
                rows.append(row)
            cand.append({"name": a[1:].strip(), "hdr": hdr, "rows": rows})
    out["cand"] = cand

    # 시장분석 (스왑용) — 헤더명으로 열 찾기
    M = wb["시장분석"]
    heads = [str(M.cell(1, c).value or "") for c in range(1, M.max_column + 1)]

    def col(prefix):
        for i, h in enumerate(heads):
            if h.split("\n")[0] == prefix:
                return i + 1
        raise KeyError(prefix)

    want = {"name": "채권명", "iss": "발행자", "sec": "섹터", "reg": "지역", "rtg": "등급",
            "dur": "듀레이션", "ytm": "YTM(%)", "g": "G-spread", "ery": "1Y기대수익률(%)",
            "be": "손익분기 확대폭(bp)", "res": "잔차(bp)", "z": "z", "crv": "적용커브",
            "hmv": "보유 평가금액($M,선택)", "chk": "제약점검", "rv": "RV대상"}
    cidx = {k: col(v) for k, v in want.items()}
    mkt = []
    for r in range(2, M.max_row + 1):
        nm = M.cell(r, cidx["name"]).value
        if not nm:
            continue
        rec = {k: M.cell(r, c).value for k, c in cidx.items()}
        for k in ("rtg", "chk", "crv", "sec", "reg", "iss"):
            if rec[k] in (None, 0):
                rec[k] = ""
        mkt.append(rec)
    out["mkt"] = mkt

    W = wb["스왑분석"]
    out["swap_default"] = {"name": _v(W, "C4"), "tol": _v(W, "C5") or 1.0,
                           "ssec": _v(W, "C6") or "N", "sreg": _v(W, "C7") or "N",
                           "xheld": _v(W, "C8") or "N", "minp": _v(W, "C9") or 0,
                           "amt": _v(W, "C10")}
    return out


# ---------------------------------------------------------------- 렌더
CSS = r"""
#port-tab{--ink:var(--tx,#14213d);--muted:var(--tx2,#5c677d);--faint:var(--tx3,#a3abb9);--line:var(--bd,#d9dee7);
 --field:var(--sf2,#ffffff);--soft:var(--sf2,#f3f5f9);--accent:var(--ac,#2f4b8a);
 --cheap:var(--gn,#1d7a46);--cheap-bg:rgba(16,185,129,.14);--rich:var(--rd,#b3261e);--rich-bg:rgba(239,68,68,.14);
 color:var(--ink);font-variant-numeric:tabular-nums;font-size:13px;line-height:1.45}
#port-tab *{box-sizing:border-box}
#port-tab h2{font-size:20px;margin:0 0 2px;font-weight:700;letter-spacing:-.01em}
#port-tab h3{font-size:15px;margin:0 0 8px;font-weight:700}
#port-tab .pt-sub{color:var(--muted);margin:0 0 14px}
#port-tab .pt-warn{color:var(--rich);font-weight:600;margin:0 0 10px}
#port-tab .pt-grid{display:grid;gap:18px;grid-template-columns:minmax(260px,330px) 1fr;align-items:start;margin-bottom:26px}
#port-tab .pt-kpi{border-top:3px solid var(--ink)}
#port-tab .pt-kpi div{display:flex;justify-content:space-between;gap:12px;padding:5px 0;border-bottom:1px solid var(--line)}
#port-tab .pt-kpi div span:last-child{font-weight:600;text-align:right;white-space:nowrap}
#port-tab .pt-kpi div.pt-lead span:last-child{font-size:17px}
#port-tab section.pt-block{margin-bottom:30px}
#port-tab .pt-scroll{overflow-x:auto;-webkit-overflow-scrolling:touch}
#port-tab table{border-collapse:collapse;width:100%;white-space:nowrap}
#port-tab th{background:var(--soft);color:var(--muted);font-weight:600;font-size:12px;text-align:right;padding:6px 8px;border-bottom:1px solid var(--line);vertical-align:bottom;white-space:normal;min-width:64px;position:static;cursor:default;text-transform:none;letter-spacing:0}
#port-tab th:first-child,#port-tab td:first-child{text-align:left}
#port-tab td{padding:5px 8px;border-bottom:1px solid var(--line);text-align:right;font-size:12px}
#port-tab td.t,#port-tab th.t{text-align:left}
#port-tab .pt-grid>*,#port-tab .pt-two>*{min-width:0}
#port-tab tr.pt-total td{font-weight:700;border-top:2px solid var(--ink)}
#port-tab tr.pt-sel td{background:var(--soft);font-weight:600}
#port-tab .pos{color:var(--cheap)} #port-tab .neg{color:var(--rich)}
#port-tab .pt-cheap{background:var(--cheap-bg);color:var(--cheap);font-weight:600}
#port-tab .pt-rich{background:var(--rich-bg);color:var(--rich);font-weight:600}
#port-tab .pt-dim{color:var(--faint)}
#port-tab .pt-tabs{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 10px}
#port-tab .pt-tabs button{border:1px solid var(--line);background:transparent;color:var(--ink);padding:4px 10px;border-radius:14px;cursor:pointer;font:inherit}
#port-tab .pt-tabs button[aria-pressed="true"]{background:var(--accent);color:#fff;border-color:var(--accent)}
#port-tab .pt-two{display:grid;gap:18px;grid-template-columns:1fr 1fr}
#port-tab .pt-ctl{display:flex;flex-wrap:wrap;gap:10px 16px;align-items:end;margin-bottom:12px}
#port-tab .pt-ctl label{display:flex;flex-direction:column;gap:3px;color:var(--muted);font-size:12px}
#port-tab .pt-ctl select,#port-tab .pt-ctl input{font:inherit;color:var(--ink);padding:4px 6px;border:1px solid var(--line);border-radius:4px;background:var(--field)}
#port-tab .pt-ctl select#pt-sw-name{min-width:240px}
#port-tab .pt-ctl input[type=number]{width:84px}
#port-tab .pt-note{color:var(--muted);font-size:12px;margin-top:6px}
#port-tab button:focus-visible,#port-tab select:focus-visible,#port-tab input:focus-visible{outline:2px solid var(--accent);outline-offset:1px}
@media (max-width:900px){#port-tab .pt-grid,#port-tab .pt-two{grid-template-columns:1fr}}
"""

JS = r"""
(function(){
const D=JSON.parse(document.getElementById('pt-data').textContent);
const $=s=>document.querySelector('#port-tab '+s);
const f=(v,d=1,sign=false)=>{if(v===null||v===undefined||v===''||isNaN(v))return '';const s=Number(v).toLocaleString('en-US',{minimumFractionDigits:d,maximumFractionDigits:d});return sign&&v>0?'+'+s:s};
const cls=v=>v>0?'pos':(v<0?'neg':'');
// ---------- bucket tabs
const BCOLS=[[1,'평가금액($M)',v=>f(v,1)],[2,'비중',v=>f(v*100,1)+'%'],[4,'DV01비중',v=>f(v*100,1)+'%'],[5,'가중YTM(%)',v=>f(v,3)],
 [6,'포트대비(bp)',v=>f(v,1,true),true],[7,'G-spread(bp)',v=>f(v,1)],[8,'듀레이션',v=>f(v,2)],[9,'손익분기 확대폭(bp)',v=>f(v,1)],
 [10,'1Y기대수익률',v=>v===''||v==null?'':f(v*100,2)+'%'],[12,'Cheap 보유',v=>f(v,0)],[13,'Rich 보유',v=>f(v,0)],
 [14,'목표비중',v=>v==null||v===''?'':f(v*100,1)+'%'],[15,'갭',v=>v==null||v===''?'':f(v*100,1,true)+'%',true],[17,'쇼크손실 비중',v=>f(v*100,1)+'%']];
function drawBucket(i){
  const b=D.buckets[i];
  document.querySelectorAll('#port-tab .pt-tabs button').forEach((x,j)=>x.setAttribute('aria-pressed',j===i));
  let h='<table><thead><tr><th>구분</th>'+BCOLS.map(c=>'<th>'+c[1]+'</th>').join('')+'</tr></thead><tbody>';
  b.rows.forEach(r=>{
    const empty=!r[1]&&r[0]!=='합계'; if(empty) return;
    h+='<tr class="'+(r[0]==='합계'?'pt-total':'')+'"><td>'+r[0]+'</td>'+BCOLS.map(c=>{const v=r[c[0]];return '<td class="'+(c[3]&&typeof v==='number'?cls(v):'')+'">'+c[2](v)+'</td>'}).join('')+'</tr>';
  });
  $('#pt-bucket').innerHTML=h+'</tbody></table>';
}
const tabs=$('.pt-tabs');
D.buckets.forEach((b,i)=>{const bt=document.createElement('button');bt.type='button';bt.textContent=b.name.replace(/\s*\(.*\)/,'');bt.onclick=()=>drawBucket(i);tabs.appendChild(bt)});
if(D.buckets.length) drawBucket(0);
// ---------- swap
const M=D.mkt; const byName={}; M.forEach((m,i)=>byName[m.name]=i);
const sel=$('#pt-sw-name');
const held=M.filter(m=>m.hmv>0).sort((a,b)=>b.hmv-a.hmv), rest=M.filter(m=>!(m.hmv>0)).sort((a,b)=>a.name.localeCompare(b.name));
[['보유 종목',held],['기타 시장 종목',rest]].forEach(([lab,arr])=>{const g=document.createElement('optgroup');g.label=lab;arr.forEach(m=>{const o=document.createElement('option');o.value=m.name;o.textContent=m.name+(m.hmv>0?'  ('+f(m.hmv,1)+'M)':'');g.appendChild(o)});sel.appendChild(g)});
const d=D.swap_default; sel.value=d.name; $('#pt-sw-tol').value=d.tol; $('#pt-sw-sec').checked=d.ssec==='Y'; $('#pt-sw-reg').checked=d.sreg==='Y'; $('#pt-sw-held').checked=d.xheld==='Y'; $('#pt-sw-min').value=d.minp; if(d.amt) $('#pt-sw-amt').value=d.amt;
function drawSwap(){
  const s=M[byName[sel.value]]; if(!s){return}
  const tol=+$('#pt-sw-tol').value||1, minp=+$('#pt-sw-min').value||0;
  const amtIn=parseFloat($('#pt-sw-amt').value); const amt=!isNaN(amtIn)?amtIn:(s.hmv>0?s.hmv:10);
  $('#pt-sw-amt-eff').textContent='적용 매도금액 '+f(amt,2)+'M'+(isNaN(amtIn)&&!(s.hmv>0)?' (보유분 없음 → 10M 가정)':'');
  const sc=[['채권명',s.name],['섹터',s.sec],['지역',s.reg],['등급',s.rtg],['듀레이션',f(s.dur,2)],['YTM(%)',f(s.ytm,3)],['G-spread',f(s.g,1)],['1Y기대수익률(%)',f(s.ery,3)],['손익분기(bp)',f(s.be,1)],['잔차(bp)',f(s.res,1,true)],['z',f(s.z,2,true)],['보유($M)',f(s.hmv,2)]];
  $('#pt-sw-sell').innerHTML='<table><thead><tr>'+sc.map((x,i)=>'<th'+(i<4?' class="t"':'')+'>'+x[0]+'</th>').join('')+'</tr></thead><tbody><tr class="pt-sel">'+sc.map((x,i)=>'<td'+(i<4?' class="t"':'')+'>'+x[1]+'</td>').join('')+'</tr></tbody></table>';
  const ok=m=>m.rv===1&&m.name!==s.name&&typeof m.dur==='number'&&Math.abs(m.dur-s.dur)<=tol&&!/FRN|BBB0미만/.test(m.chk||'')
    &&(!$('#pt-sw-sec').checked||m.sec===s.sec)&&(!$('#pt-sw-reg').checked||m.reg===s.reg)&&(!$('#pt-sw-held').checked||!(m.hmv>0))&&(m.g-s.g)>=minp;
  const alts=M.filter(ok).sort((a,b)=>(b.g-s.g)-(a.g-s.g)).slice(0,20);
  const H=['채권명','섹터','지역','등급','듀레이션','G-spread','G 픽업','YTM 픽업','기대수익률 픽업','손익분기 변화','잔차','z','피어','DV01중립 매수($M)','연간 기대수익 변화($)','DV01 변화($/bp)','보유($M)'];
  let h='<table><thead><tr><th>#</th>'+H.map((x,i)=>'<th'+(i<4||x==='피어'?' class="t"':'')+'>'+x+'</th>').join('')+'</tr></thead><tbody>';
  alts.forEach((m,i)=>{
    const same=m.crv===s.crv, gp=m.g-s.g, yp=(m.ytm-s.ytm)*100, ep=(m.ery-s.ery)*100, bd=(m.be!==''&&s.be!=='')?m.be-s.be:null;
    const dim=same?'':' pt-dim';
    h+='<tr><td>'+(i+1)+'</td><td class="t">'+m.name+'</td><td class="t">'+m.sec+'</td><td class="t">'+m.reg+'</td><td class="t">'+m.rtg+'</td><td>'+f(m.dur,2)+'</td><td>'+f(m.g,1)+'</td>'
     +'<td class="'+cls(gp)+'">'+f(gp,1,true)+'</td><td class="'+cls(yp)+'">'+f(yp,1,true)+'</td><td class="'+cls(ep)+'">'+f(ep,1,true)+'</td><td class="'+cls(bd)+'">'+f(bd,1,true)+'</td>'
     +'<td class="'+dim+'">'+f(m.res,1,true)+'</td><td class="'+dim+'">'+f(m.z,2,true)+'</td><td class="t'+dim+'">'+(same?'동일 피어':'다른 피어')+'</td>'
     +'<td>'+f(amt*s.dur/m.dur,2)+'</td><td class="'+cls(ep)+'">'+f(amt*ep/10000*1e6,0,true)+'</td><td>'+f(amt*1e6*(m.dur-s.dur)/1e4,0,true)+'</td><td>'+f(m.hmv,2)+'</td></tr>';
  });
  if(!alts.length) h+='<tr><td colspan="18" class="t">조건에 맞는 대안이 없습니다. 듀레이션 허용범위를 넓히거나 필터를 해제하세요.</td></tr>';
  $('#pt-sw-alt').innerHTML=h+'</tbody></table>';
}
['#pt-sw-name','#pt-sw-tol','#pt-sw-sec','#pt-sw-reg','#pt-sw-held','#pt-sw-min','#pt-sw-amt'].forEach(id=>{$(id).addEventListener('input',drawSwap);$(id).addEventListener('change',drawSwap)});
drawSwap();
})();
"""


def _fmt(v, kind):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return html.escape(v)
    if kind == "pct":
        return f"{v*100:,.2f}%"
    if kind == "pct1":
        return f"{v*100:,.1f}%"
    if kind == "int":
        return f"{v:,.0f}"
    if kind == "bp":
        return f"{v:,.1f}"
    if kind == "sbp":
        return f"{v:+,.1f}"
    if kind == "sz":
        return f"{v:+.2f}"
    if kind == "2":
        return f"{v:,.2f}"
    if kind == "3":
        return f"{v:,.3f}"
    return f"{v:,.1f}"


KPI_FMT = {"평가금액": "1", "액면": "1", "가중 YTM": "3", "가중 G-spread": "bp", "가중 듀레이션": "2",
           "DV01": "int", "CR01": "int", "1Y 캐리": "int", "1Y 스프레드": "int", "1Y 기대수익($)": "int",
           "1Y 기대수익률": "pct", "평가손익": "int", "국채 비중": "pct1", "쇼크조정": "int", "G-spread 커버리지": "pct1"}


def _kpi_fmt(k, v):
    for key, kind in KPI_FMT.items():
        if k.startswith(key):
            return _fmt(v, kind)
    return _fmt(v, "1")


def _signcls(v):
    return "pos" if isinstance(v, (int, float)) and v > 0 else ("neg" if isinstance(v, (int, float)) and v < 0 else "")


def render_port_tab(xlsx_path, data=None):
    d = data or read_workbook(xlsx_path)
    e = html.escape
    lead = {"평가금액($M)", "가중 YTM(%)", "1Y 기대수익률(%, 평가금액 대비)"}
    kpi_html = "".join(
        f'<div class="{"pt-lead" if x["k"] in lead else ""}"><span>{e(x["k"])}</span><span>{_kpi_fmt(x["k"], x["v"])}</span></div>'
        for x in d["kpi"])

    sh = d["scen"]["hdr"]
    scen_rows = ""
    for r in d["scen"]["rows"]:
        scen_rows += (f'<tr><td>{e(str(r[0]))}</td><td>{_fmt(r[1],"int")}</td><td>{_fmt(r[2],"int")}</td>'
                      f'<td>{_fmt(r[3],"int")}</td><td class="{_signcls(r[4])}">{_fmt(r[4],"int")}</td>'
                      f'<td class="{_signcls(r[6])}">{_fmt(r[6],"int")}</td><td>{_fmt(r[7],"int")}</td>'
                      f'<td class="{_signcls(r[8])}"><b>{_fmt(r[8],"pct")}</b></td></tr>')
    scen_html = (f'<table><thead><tr><th>시나리오</th><th>금리Δ(bp)</th><th>스프레드Δ(bp)</th><th>캐리+롤다운($)</th>'
                 f'<th>금리손익($)</th><th>스프레드손익($, 차등·비례)</th><th>총수익($)</th><th>총수익률</th></tr></thead>'
                 f'<tbody>{scen_rows}</tbody></table>')
    be_html = "".join(f'<div><span>{e(str(x["k"]))}</span><span>{_fmt(x["v"],"bp")}bp</span></div>' for x in d["breakeven"])

    fund_rows = "".join(
        f'<tr><td>{e(str(r[0]))}</td><td>{_fmt(r[1],"1")}</td><td>{_fmt(r[2],"3")}</td><td>{_fmt(r[3],"bp")}</td>'
        f'<td>{_fmt(r[4],"2")}</td><td>{_fmt(r[5],"int")}</td><td>{_fmt(r[6],"pct")}</td><td>{_fmt(r[7],"int")}</td></tr>'
        for r in d["funds"]["rows"])
    fund_html = ('<table><thead><tr><th>운용코드</th><th>평가금액($M)</th><th>가중YTM(%)</th><th>가중G(bp)</th>'
                 '<th>듀레이션</th><th>DV01($/bp)</th><th>1Y기대수익률</th><th>종목수</th></tr></thead>'
                 f'<tbody>{fund_rows}</tbody></table>')

    def cand_table(c):
        rows = ""
        for r in c["rows"]:
            z = r[11]
            zc = "pt-cheap" if isinstance(z, (int, float)) and z > 0 else "pt-rich"
            rows += (f'<tr><td>{e(str(r[1]))}</td><td class="t">{e(str(r[3] or ""))}</td><td class="t">{e(str(r[4] or ""))}</td>'
                     f'<td class="t">{e(str(r[5] or ""))}</td><td>{_fmt(r[6],"2")}</td><td>{_fmt(r[8],"bp")}</td>'
                     f'<td>{_fmt(r[9],"bp")}</td><td class="{_signcls(r[10])}">{_fmt(r[10],"sbp")}</td>'
                     f'<td class="{zc}">{_fmt(z,"sz")}</td><td>{_fmt(r[12],"sbp")}</td><td>{_fmt(r[13],"2")}</td>'
                     f'<td class="t">{e(str(r[14] or ""))}</td></tr>')
        return (f'<div><h3>{e(c["name"])}</h3><div class="pt-scroll"><table><thead><tr><th>채권명</th><th class="t">국가</th><th class="t">섹터</th>'
                f'<th class="t">등급</th><th>듀레이션</th><th>G</th><th>적정G</th><th>잔차</th><th>z</th><th>발행자내</th>'
                f'<th>보유($M)</th><th class="t">점검</th></tr></thead><tbody>{rows}</tbody></table></div></div>')

    cand_html = "".join(cand_table(c) for c in d["cand"])
    payload = {"buckets": d["buckets"], "mkt": d["mkt"], "swap_default": d["swap_default"]}
    data_json = json.dumps(payload, ensure_ascii=False, default=str).replace("</", "<\\/")
    warn = f'<p class="pt-warn">{e(d["warning"])}</p>' if d["warning"] else ""

    return f"""<section id="port-tab">
<style>{CSS}</style>
<h2>포트 운용</h2>
<p class="pt-sub">{e(d["label"])} · 엑셀 기준값, 생성 {d["generated"]}</p>
{warn}
<div class="pt-grid">
  <div class="pt-kpi">{kpi_html}</div>
  <div>
    <h3>1년 보유 시나리오</h3><div class="pt-scroll">{scen_html}</div>
    <div class="pt-kpi" style="margin-top:12px;border-top-width:1px">{be_html}</div>
    <h3 style="margin-top:18px">운용코드별</h3><div class="pt-scroll">{fund_html}</div>
  </div>
</div>
<section class="pt-block">
  <h3>버킷분석</h3>
  <div class="pt-tabs" role="group" aria-label="버킷 구분"></div>
  <div class="pt-scroll" id="pt-bucket"></div>
  <p class="pt-note">손익분기 확대폭 = (G-spread + 스프레드 롤다운) ÷ 크레딧 듀레이션. 쇼크손실 비중은 차등·비례 스프레드 쇼크 기준.</p>
</section>
<section class="pt-block">
  <div class="pt-two">{cand_html}</div>
  <p class="pt-note">z는 각 종목의 피어커브(지역·섹터) 대비 위치. 다른 피어 종목끼리 잔차 크기를 비교하지 말 것.</p>
</section>
<section class="pt-block">
  <h3>스왑분석</h3>
  <div class="pt-ctl">
    <label>매도 종목<select id="pt-sw-name"></select></label>
    <label>듀레이션 ±(년)<input id="pt-sw-tol" type="number" step="0.25" min="0"></label>
    <label>최소 G 픽업(bp)<input id="pt-sw-min" type="number" step="5"></label>
    <label>매도금액($M)<input id="pt-sw-amt" type="number" step="0.5" placeholder="보유분"></label>
    <label style="flex-direction:row;align-items:center;gap:5px"><input id="pt-sw-sec" type="checkbox">같은 섹터만</label>
    <label style="flex-direction:row;align-items:center;gap:5px"><input id="pt-sw-reg" type="checkbox">같은 지역만</label>
    <label style="flex-direction:row;align-items:center;gap:5px"><input id="pt-sw-held" type="checkbox">보유 종목 제외</label>
  </div>
  <p class="pt-note" id="pt-sw-amt-eff"></p>
  <div class="pt-scroll" id="pt-sw-sell"></div>
  <div class="pt-scroll" id="pt-sw-alt" style="margin-top:10px"></div>
  <p class="pt-note">픽업 = 대안 − 매도 종목. 연간 기대수익 변화·DV01 변화는 같은 금액으로 갈아탈 때 기준. 다른 피어의 잔차·z는 흐리게 표시(직접 비교 불가). FRN·BBB0 미만은 대안에서 제외.</p>
</section>
<script type="application/json" id="pt-data">{data_json}</script>
<script>{JS}</script>
</section>"""


def main():
    if len(sys.argv) < 2:
        print("usage: python export_port_tab.py <포트운용 xlsx> [출력폴더]")
        sys.exit(1)
    xlsx = Path(sys.argv[1]); outdir = Path(sys.argv[2]) if len(sys.argv) > 2 else xlsx.parent
    data = read_workbook(xlsx)
    frag = render_port_tab(xlsx, data)
    (outdir / "port_tab.json").write_text(json.dumps(data, ensure_ascii=False, default=str, indent=1), encoding="utf-8")
    page = ("<!doctype html><html lang='ko'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width, initial-scale=1, viewport-fit=cover'>"
            "<title>포트 운용</title><style>body{margin:0;padding:24px;font-family:'Pretendard','Apple SD Gothic Neo','Malgun Gothic',system-ui,sans-serif;background:#fff}</style>"
            f"</head><body>{frag}</body></html>")
    (outdir / "port_tab_preview.html").write_text(page, encoding="utf-8")
    print("written:", outdir / "port_tab_preview.html", outdir / "port_tab.json")


if __name__ == "__main__":
    main()
